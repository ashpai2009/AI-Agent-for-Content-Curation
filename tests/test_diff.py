"""Diff fidelity and reconciliation.

Two obligations, and the second is the one that matters. The diff must report *nothing*
when nothing changed, or every job fails its gate for no reason. And it must report
*something* for a change in each compared dimension, or the gate is decoration.
"""

from __future__ import annotations

import shutil
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from conftest import write_feature_rich_workbook
from oatutor_council.models import ChangeRecord, ColumnKey
from oatutor_council.workbook.diff import (
    DiffDimension,
    compare_workbooks,
    reconcile,
)


def resaved(source: Path, target: Path) -> Path:
    """A byte copy put through one openpyxl load-and-save, exactly as a job does."""
    shutil.copy2(source, target)
    workbook = load_workbook(target, data_only=False)
    workbook.save(target)
    workbook.close()
    return target


def change(row: int, column: int, before: str, after: str, **kwargs) -> ChangeRecord:
    return ChangeRecord(
        change_id=f"c{row}-{column}",
        issue_id="issue-1",
        patch_id="patch-1",
        block_id="block-0000",
        row=row,
        column=column,
        before=before,
        after=after,
        applied_at=datetime.now(timezone.utc),
        **kwargs,
    )


# --------------------------------------------------------------------------------------
# Fidelity
# --------------------------------------------------------------------------------------


def test_a_zero_edit_resave_produces_an_empty_diff(feature_rich_workbook, tmp_path):
    """The load-bearing test for the no-escape-hatch rule. If openpyxl perturbed
    anything here, the honest fix is a narrow allowlist entry naming that exact property
    -- never a general 'formatting normalisation' category, which would absorb real
    damage alongside it."""
    output = resaved(feature_rich_workbook, tmp_path / "out.xlsx")
    differences = compare_workbooks(feature_rich_workbook, output)
    assert differences == (), [d.describe() for d in differences]


def test_an_empty_string_cell_and_an_empty_cell_compare_equal(tmp_path):
    """The second and last tolerated normalisation. openpyxl rewrites a cell holding an
    empty string as a genuinely empty cell, changing its type from `inlineStr` to `n`.
    Both render as "", so nothing is lost -- but the exception is scoped to cells empty
    on both sides, and the second half of this test is what keeps it honest."""
    from openpyxl import Workbook

    source = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = ""
    workbook.active["A2"] = "content"
    workbook.save(source)
    workbook.close()

    resave = resaved(source, tmp_path / "b.xlsx")
    assert compare_workbooks(source, resave) == ()

    emptied = tmp_path / "c.xlsx"
    shutil.copy2(source, emptied)
    workbook = load_workbook(emptied)
    workbook.active["A2"] = None  # real content lost, not a normalisation
    workbook.save(emptied)
    workbook.close()
    dimensions = {d.dimension for d in compare_workbooks(source, emptied)}
    assert DiffDimension.CELL_VALUE in dimensions
    assert DiffDimension.CELL_TYPE in dimensions


def test_shrink_to_fit_none_and_false_are_the_same_thing(tmp_path):
    """The single normalisation observed on the real corpus. In OOXML an absent
    attribute means false, so `None` and `False` must compare equal -- while a genuine
    `True -> None` must still be reported, which the next assertion pins."""
    from openpyxl import Workbook

    source = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active["A1"].alignment = Alignment(shrink_to_fit=False)
    workbook.save(source)
    workbook.close()

    output = tmp_path / "b.xlsx"
    workbook = Workbook()
    workbook.active["A1"].alignment = Alignment(shrink_to_fit=None)
    workbook.save(output)
    workbook.close()
    assert compare_workbooks(source, output) == ()

    lost = tmp_path / "c.xlsx"
    workbook = Workbook()
    workbook.active["A1"].alignment = Alignment(shrink_to_fit=True)
    workbook.save(lost)
    workbook.close()
    differences = compare_workbooks(lost, output)
    assert [d.prop for d in differences] == ["shrink_to_fit"]


# --------------------------------------------------------------------------------------
# Every dimension is actually compared
# --------------------------------------------------------------------------------------


def mutate(path: Path, mutation) -> Path:
    workbook = load_workbook(path, data_only=False)
    mutation(workbook, workbook["Main"])
    workbook.save(path)
    workbook.close()
    return path


MUTATIONS = {
    DiffDimension.CELL_VALUE: lambda wb, ws: ws.__setitem__("A1", "tampered"),
    DiffDimension.CELL_TYPE: lambda wb, ws: ws.__setitem__("B1", "not a number"),
    DiffDimension.NUMBER_FORMAT: lambda wb, ws: setattr(
        ws["B1"], "number_format", "0.00%"
    ),
    DiffDimension.FONT: lambda wb, ws: setattr(ws["A1"], "font", Font(name="Courier")),
    DiffDimension.FILL: lambda wb, ws: setattr(
        ws["A1"], "fill", PatternFill(fill_type="solid", fgColor="00FF00")
    ),
    DiffDimension.BORDER: lambda wb, ws: setattr(
        ws["A1"], "border", Border(left=Side(style="thick"))
    ),
    DiffDimension.ALIGNMENT: lambda wb, ws: setattr(
        ws["A1"], "alignment", Alignment(horizontal="right")
    ),
    DiffDimension.ROW_HEIGHT: lambda wb, ws: setattr(ws.row_dimensions[1], "height", 99),
    DiffDimension.ROW_HIDDEN: lambda wb, ws: setattr(
        ws.row_dimensions[5], "hidden", False
    ),
    DiffDimension.COLUMN_WIDTH: lambda wb, ws: setattr(
        ws.column_dimensions["A"], "width", 8
    ),
    DiffDimension.COLUMN_HIDDEN: lambda wb, ws: setattr(
        ws.column_dimensions["E"], "hidden", False
    ),
    DiffDimension.MERGED_RANGES: lambda wb, ws: ws.unmerge_cells("A3:C4"),
    DiffDimension.FREEZE_PANES: lambda wb, ws: setattr(ws, "freeze_panes", None),
    DiffDimension.SHEET_SET: lambda wb, ws: wb.remove(wb["Second"]),
    DiffDimension.SHEET_ORDER: lambda wb, ws: wb.move_sheet("Second", -1),
    DiffDimension.SHEET_VISIBILITY: lambda wb, ws: setattr(
        wb["Hidden"], "sheet_state", "visible"
    ),
    DiffDimension.HYPERLINK: lambda wb, ws: setattr(ws["G1"], "hyperlink", None),
    DiffDimension.DATA_VALIDATION: lambda wb, ws: ws.data_validations.dataValidation.clear(),
    DiffDimension.IMAGE: lambda wb, ws: ws._images.clear(),
}


def test_every_compared_dimension_has_a_mutation_test():
    """A dimension added to the enum without a test here is a gate with a hole in it."""
    assert set(MUTATIONS) == set(DiffDimension)


@pytest.mark.parametrize("dimension", sorted(MUTATIONS, key=str))
def test_a_change_in_each_dimension_is_detected(
    feature_rich_workbook, tmp_path, dimension
):
    output = mutate(
        resaved(feature_rich_workbook, tmp_path / "out.xlsx"), MUTATIONS[dimension]
    )
    found = {d.dimension for d in compare_workbooks(feature_rich_workbook, output)}
    assert dimension in found, f"{dimension} went unnoticed; found {sorted(found)}"


# --------------------------------------------------------------------------------------
# Reconciliation
# --------------------------------------------------------------------------------------


def test_an_edit_in_the_ledger_explains_its_own_value_change(
    feature_rich_workbook, tmp_path
):
    output = mutate(
        resaved(feature_rich_workbook, tmp_path / "out.xlsx"),
        lambda wb, ws: ws.__setitem__("A1", "corrected"),
    )
    differences = compare_workbooks(feature_rich_workbook, output)
    unexplained = reconcile(
        differences, [change(1, 1, "value", "corrected")], sheet_name="Main"
    )
    assert unexplained == ()


def test_an_edit_with_the_wrong_after_value_stays_unexplained(
    feature_rich_workbook, tmp_path
):
    """A ledger entry is not a licence to change the cell however you like. The recorded
    `after` must be what is actually on disk, or the change log and the workbook have
    quietly diverged."""
    output = mutate(
        resaved(feature_rich_workbook, tmp_path / "out.xlsx"),
        lambda wb, ws: ws.__setitem__("A1", "something else entirely"),
    )
    differences = compare_workbooks(feature_rich_workbook, output)
    unexplained = reconcile(
        differences, [change(1, 1, "value", "corrected")], sheet_name="Main"
    )
    assert [d.dimension for d in unexplained] == [DiffDimension.CELL_VALUE]


def test_an_unauthorised_change_elsewhere_survives_reconciliation(
    feature_rich_workbook, tmp_path
):
    """The case the gate exists for: one authorised edit, plus a merged range quietly
    lost. A value-only diff would pass this workbook."""

    def tamper(wb, ws):
        ws["A1"] = "corrected"
        ws.unmerge_cells("A3:C4")

    output = mutate(resaved(feature_rich_workbook, tmp_path / "out.xlsx"), tamper)
    unexplained = reconcile(
        compare_workbooks(feature_rich_workbook, output),
        [change(1, 1, "value", "corrected")],
        sheet_name="Main",
    )
    assert [d.dimension for d in unexplained] == [DiffDimension.MERGED_RANGES]


def test_the_edited_row_rule_is_forgiven_only_on_edited_rows(
    feature_rich_workbook, tmp_path
):
    def edit_row_one(wb, ws):
        ws["A1"] = "corrected"
        ws["A1"].alignment = Alignment(
            horizontal="center", vertical="top", wrap_text=False, indent=2,
            text_rotation=45,
        )
        ws.row_dimensions[1].height = 15.0

    output = mutate(resaved(feature_rich_workbook, tmp_path / "out.xlsx"), edit_row_one)
    unexplained = reconcile(
        compare_workbooks(feature_rich_workbook, output),
        [change(1, 1, "value", "corrected")],
        sheet_name="Main",
    )
    assert unexplained == ()


def test_wrap_removal_on_an_untouched_row_is_not_forgiven(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "a.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "one"
    sheet["A2"] = "two"
    for ref in ("A1", "A2"):
        sheet[ref].alignment = Alignment(wrap_text=True)
    workbook.save(source)
    workbook.close()

    output = tmp_path / "b.xlsx"
    shutil.copy2(source, output)
    workbook = load_workbook(output)
    workbook.active["A1"] = "corrected"
    workbook.active["A1"].alignment = Alignment(wrap_text=False)
    workbook.active["A2"].alignment = Alignment(wrap_text=False)  # never edited
    workbook.save(output)
    workbook.close()

    unexplained = reconcile(
        compare_workbooks(source, output),
        [change(1, 1, "one", "corrected")],
        sheet_name="Sheet",
    )
    assert [(d.dimension, d.row) for d in unexplained] == [
        (DiffDimension.ALIGNMENT, 2)
    ]


def test_the_text_number_format_is_forgiven_only_on_answer_and_choice_columns(tmp_path):
    """Forcing `@` is the defence against Excel turning a repaired fraction back into a
    date. It is authorised on answer and choice cells and nowhere else."""
    from openpyxl import Workbook

    source = tmp_path / "a.xlsx"
    workbook = Workbook()
    workbook.active["E2"] = "x"
    workbook.active["C2"] = "y"
    workbook.save(source)
    workbook.close()

    output = tmp_path / "b.xlsx"
    shutil.copy2(source, output)
    workbook = load_workbook(output)
    for ref in ("E2", "C2"):
        workbook.active[ref] = "corrected"
        workbook.active[ref].number_format = "@"
    workbook.save(output)
    workbook.close()

    unexplained = reconcile(
        compare_workbooks(source, output),
        [
            change(2, 5, "x", "corrected", column_key=ColumnKey.ANSWER),
            change(2, 3, "y", "corrected", column_key=ColumnKey.TITLE),
        ],
        sheet_name="Sheet",
    )
    assert [(d.dimension, d.column) for d in unexplained] == [
        (DiffDimension.NUMBER_FORMAT, 3)
    ]


def test_a_difference_on_another_sheet_is_never_explained(
    feature_rich_workbook, tmp_path
):
    """The ledger only ever describes the curated sheet, so a change on a different
    sheet has nothing that could authorise it."""
    output = mutate(
        resaved(feature_rich_workbook, tmp_path / "out.xlsx"),
        lambda wb, ws: wb["Second"].__setitem__("A1", "injected"),
    )
    unexplained = reconcile(
        compare_workbooks(feature_rich_workbook, output),
        [change(1, 1, "value", "corrected")],
        sheet_name="Main",
    )
    # Value and type both, since the cell went from empty to text.
    assert {d.sheet for d in unexplained} == {"Second"}
    assert {d.dimension for d in unexplained} == {
        DiffDimension.CELL_VALUE,
        DiffDimension.CELL_TYPE,
    }
