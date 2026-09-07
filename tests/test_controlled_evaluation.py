"""The held-out scorer accepts invariants, not one synthetic author's favorite text."""

from __future__ import annotations

from openpyxl import load_workbook

from conftest import problem, step
from scripts.evaluate_controlled_run import (
    allowed_related_cells,
    evaluate_item,
    keyed_items,
    supplementary_items,
)


def test_multiple_valid_mc_repairs_are_scored_by_the_grading_contract(make_workbook):
    path = make_workbook(
        [
            problem("p1", title="Choose the exact probability", oer_src="s", license="CC"),
            step(
                "p1",
                answer="1/2",
                answer_type="mc",
                mc_choices="1/2|1/3|1/4|3/4",
            ),
        ]
    )
    workbook = load_workbook(path, data_only=False)
    try:
        passed, _, _ = evaluate_item(
            {"cell": "I3", "predicate": "mc_exact_answer_once"}, workbook.active
        )
    finally:
        workbook.close()
    assert passed


def test_equivalent_ordered_pair_distractor_fails_the_mc_contract(make_workbook):
    path = make_workbook(
        [
            problem("p1", title="Choose the pair", oer_src="s", license="CC"),
            step(
                "p1",
                answer="(3,2)",
                answer_type="mc",
                mc_choices="(3,2)|(2,3)|(6/2,2)|(4,1)",
            ),
        ]
    )
    workbook = load_workbook(path, data_only=False)
    try:
        passed, _, _ = evaluate_item(
            {"cell": "I3", "predicate": "mc_exact_answer_once"}, workbook.active
        )
    finally:
        workbook.close()
    assert not passed


def test_an_equivalent_graded_answer_is_not_forced_to_one_string(make_workbook):
    path = make_workbook(
        [
            problem("p1", title="Solve", oer_src="s", license="CC"),
            step("p1", answer="2", answer_type="numeric"),
        ]
    )
    workbook = load_workbook(path, data_only=False)
    try:
        passed, _, _ = evaluate_item(
            {"cell": "E3", "mathEquivalentTo": "x=sqrt(4)"}, workbook.active
        )
    finally:
        workbook.close()
    assert passed


def test_a_key_without_a_machine_check_stays_manual():
    automated, manual = keyed_items(
        {
            "defectGroups": [
                {
                    "problemName": "p1",
                    "kind": "instructional_quality",
                    "corrections": [
                        {"cell": "D4", "manual": "Hint should teach without giving away the answer"}
                    ],
                }
            ]
        }
    )
    assert automated == []
    assert len(manual) == 1


def test_one_preferred_hint_sentence_is_not_an_automated_exact_check():
    """Instructional quality permits correct wording the fixture author did not choose."""
    automated, manual = keyed_items(
        {
            "defectGroups": [
                {
                    "problemName": "p1",
                    "kind": "wrong_hint",
                    "corrections": [
                        {
                            "cell": "D4",
                            "expected": "Subtract 5 from both sides, then divide by 2.",
                        }
                    ],
                }
            ]
        }
    )
    assert automated == []
    assert [item["cell"] for item in manual] == ["D4"]


def test_a_fixture_can_explicitly_require_exact_prose_when_that_is_the_contract():
    automated, manual = keyed_items(
        {
            "defectGroups": [
                {
                    "problemName": "p1",
                    "kind": "required_copy",
                    "corrections": [
                        {
                            "cell": "D4",
                            "expected": "Use this required sentence.",
                            "comparison": "exact",
                        }
                    ],
                }
            ]
        }
    )
    assert [item["cell"] for item in automated] == ["D4"]
    assert manual == []


def test_structural_cell_moves_remain_machine_checked():
    """Prose values are exact when the defect is structural, not instructional."""
    automated, manual = keyed_items(
        {
            "defectGroups": [
                {
                    "problemName": "p1",
                    "kind": "row_shift_right",
                    "corrections": [
                        {"cell": "D4", "expected": "Body text restored to its column."}
                    ],
                }
            ]
        }
    )
    assert [item["cell"] for item in automated] == ["D4"]
    assert manual == []


def test_known_source_defects_are_separate_from_the_planted_score():
    key = {
        "defectGroups": [
            {
                "problemName": "planted",
                "kind": "wrong_answer",
                "corrections": [{"cell": "E3", "expected": "2"}],
            }
        ],
        "knownSourceDefects": [
            {
                "problemName": "source",
                "kind": "rounding",
                "corrections": [{"cell": "E8", "accepted": ["0.218"]}],
            }
        ],
    }
    automated, _ = keyed_items(key)
    assert [item["cell"] for item in automated] == ["E3"]
    assert [item["cell"] for item in supplementary_items(key)] == ["E8"]


def test_coordinated_repair_cells_are_allowed_without_inflating_the_score():
    key = {
        "defectGroups": [
            {
                "problemName": "p1",
                "kind": "answer_and_choices",
                "corrections": [{"cell": "E3", "expected": "1/2"}],
                "allowedRelatedCells": ["I3"],
            }
        ]
    }
    automated, _ = keyed_items(key)
    assert [item["cell"] for item in automated] == ["E3"]
    assert allowed_related_cells(key) == {"I3"}
