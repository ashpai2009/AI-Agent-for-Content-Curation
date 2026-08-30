"""Synthetic workbook construction for the test suite.

Every fixture is generated into `tmp_path` with invented mathematics. No content from
the real curator corpus is ever copied into a committed test -- the tests reproduce the
*structures* reconnaissance found (`docs/recon-findings.md`), not the material.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Sequence

import pytest
from openpyxl import Workbook

from oatutor_council.models import FIXED_COLUMNS, ColumnKey

#: The header row as the real workbooks write it: sixteen contract columns, two blank
#: spacer columns, then the two trailing columns whose position varies in practice.
DEFAULT_HEADERS: list[str | None] = [
    "Problem Name",
    "Row Type",
    "Title",
    "Body Text",
    "Answer",
    "answerType",
    "HintID",
    "Dependency",
    "mcChoices",
    "Images (space delimited)",
    "Parent",
    "OER src",
    "openstax KC",
    "KC",
    "Taxonomy",
    "License",
    None,
    None,
    "Validator Check",
    "Time Last Checked",
]

#: Marks a deliberately blank row. `None` alone is ambiguous with "no value here", and
#: blank-row placement is the thing several segmentation tests turn on.
BLANK = object()


def cells(**values: Any) -> list[Any]:
    """Build one row from column keys, padded to the full contract width.

    Keyword names are `ColumnKey` values, e.g. `cells(problem_name="trig1",
    row_type="problem")`. Anything unset is left empty, which is what real rows do.
    """
    row: list[Any] = [None] * len(DEFAULT_HEADERS)
    for name, value in values.items():
        key = ColumnKey(name)
        index = FIXED_COLUMNS.get(key)
        if index is None:
            raise KeyError(f"{name} is not a fixed column; place it positionally")
        # An unset column is an empty cell, not a cell holding "". Writing the empty
        # string produces an `inlineStr` cell that openpyxl rewrites as empty on the
        # next save, which would show up in every diff as a spurious type change.
        row[index - 1] = value if value != "" else None
    return row


def problem(name: str, *, title: str = "", body: str = "", **extra: Any) -> list[Any]:
    return cells(
        problem_name=name, row_type="problem", title=title, body_text=body, **extra
    )


#: Row types carry a Title in 97-98% of real rows, and the rules require one. Defaulting
#: it here keeps every fixture a *well-formed* workbook whose only defects are the ones
#: the test planted -- otherwise each new rule turns every unrelated fixture into a
#: workbook with several faults and the tests start failing for reasons they are not
#: about. A test that wants the field missing passes `title=""`.
DEFAULT_TITLE = "Work through this part"
DEFAULT_BODY = "Consider what the question is asking."


def step(
    name: str,
    *,
    answer: str = "1",
    answer_type: str = "numeric",
    title: str = DEFAULT_TITLE,
    **extra: Any,
) -> list[Any]:
    return cells(
        problem_name=name,
        row_type="step",
        title=title,
        answer=answer,
        answer_type=answer_type,
        **extra,
    )


def scaffold(
    name: str,
    identifier: str,
    *,
    dependency: str = "",
    title: str = DEFAULT_TITLE,
    body_text: str = DEFAULT_BODY,
    **extra: Any,
):
    return cells(
        problem_name=name,
        row_type="scaffold",
        hint_id=identifier,
        dependency=dependency,
        title=title,
        body_text=body_text,
        **extra,
    )


def hint(
    name: str,
    identifier: str,
    *,
    body: str = "",
    title: str = DEFAULT_TITLE,
    **extra: Any,
) -> list[Any]:
    return cells(
        problem_name=name,
        row_type="hint",
        hint_id=identifier,
        title=title,
        body_text=body,
        **extra,
    )


def write_workbook(
    path: Path,
    rows: Sequence[Any],
    *,
    header_row: int = 1,
    headers: Sequence[str | None] | None = None,
    sheet_title: str = "Sheet1",
    extra_sheets: Sequence[str] = (),
) -> Path:
    """Write a synthetic workbook.

    `rows` are written consecutively from `header_row + 1`; a `BLANK` entry leaves a row
    empty. `header_row` above 1 reproduces the leading-blank-row layout one real
    workbook has, which is why the reader detects the header rather than assuming it.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    for column, label in enumerate(headers or DEFAULT_HEADERS, start=1):
        if label is not None:
            sheet.cell(row=header_row, column=column, value=label)

    for offset, row in enumerate(rows):
        target = header_row + 1 + offset
        if row is BLANK:
            continue
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=target, column=column, value=value)

    for title in extra_sheets:
        workbook.create_sheet(title=title)

    workbook.save(path)
    return path


#: A 1x1 transparent PNG, inlined so the image fixture needs no binary in the repo.
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d4944415478da6364f8cf000000030101002d0dd8a8"
    "0000000049454e44ae426082"
)


def write_feature_rich_workbook(path: Path) -> Path:
    """A workbook exercising every dimension the diff compares.

    The real corpus is almost featureless -- ten of eleven workbooks have no merged
    ranges, no explicit geometry, no validations, no images -- so round-tripping one of
    those proves very little. This fixture exists so the fidelity test has something to
    lose.
    """
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"

    sheet["A1"] = "value"
    sheet["B1"] = 3.5
    sheet["C1"] = "=B1*2"
    sheet["D1"] = True
    sheet["A1"].font = Font(
        name="Arial", size=12, bold=True, italic=True, color="FF0000"
    )
    sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    sheet["A1"].border = Border(left=Side(style="thin"), bottom=Side(style="double"))
    sheet["A1"].alignment = Alignment(
        horizontal="center", vertical="top", wrap_text=True, indent=2, text_rotation=45
    )
    sheet["B1"].number_format = "0.000"

    sheet.merge_cells("A3:C4")
    sheet.row_dimensions[1].height = 33.75
    sheet.row_dimensions[5].hidden = True
    sheet.column_dimensions["A"].width = 42.5
    sheet.column_dimensions["E"].hidden = True
    sheet.freeze_panes = "B2"

    validation = DataValidation(type="list", formula1='"a,b,c"')
    sheet.add_data_validation(validation)
    validation.add("F1:F9")

    sheet["G1"] = "link"
    sheet["G1"].hyperlink = "https://example.org/x"

    png = path.parent / f"{path.stem}-image.png"
    png.write_bytes(TINY_PNG)
    sheet.add_image(XLImage(str(png)), "H1")

    workbook.create_sheet("Second")
    workbook.create_sheet("Hidden").sheet_state = "hidden"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def feature_rich_workbook(tmp_path: Path) -> Path:
    return write_feature_rich_workbook(tmp_path / "rich.xlsx")


@pytest.fixture
def make_workbook(tmp_path: Path):
    """Return a factory writing a uniquely-named workbook into `tmp_path`."""
    counter = {"n": 0}

    def factory(rows: Sequence[Any], **kwargs: Any) -> Path:
        counter["n"] += 1
        path = tmp_path / f"workbook-{counter['n']}.xlsx"
        return write_workbook(path, rows, **kwargs)

    return factory


# --------------------------------------------------------------------------------------
# Audit coverage
# --------------------------------------------------------------------------------------
#
# A scan response must account for every graded row of the block it was sent, and a
# response that does not is re-scanned rather than accepted. A test double returning no
# coverage is therefore simulating a *non-compliant* model, and the council is right to
# reject it -- which is exactly what these helpers exist to avoid doing by accident.
#
# The rows are read back out of the payload the agent was actually given, so a double
# stays compliant when a fixture's block changes shape. That also keeps the doubles honest
# in the one way that matters here: they report on the rows they were shown, and on no
# others.

#: Row types a student answers, mirroring `ProblemBlock.graded_rows`.
GRADED_ROW_TYPES = frozenset({"step", "scaffold"})


def graded_rows_in(payload: str) -> tuple[int, ...]:
    """The graded rows of the single block rendered in this payload.

    Deliberately parses the rendered table rather than taking the block object: most
    doubles only see the request, and a helper that needed the block would have to be
    threaded through every one of them.
    """
    rows: list[int] = []
    for line in payload.splitlines():
        fields = [field.strip() for field in line.split("|")]
        if len(fields) < 3 or not fields[0].isdigit():
            continue
        if fields[2].casefold() in GRADED_ROW_TYPES:
            rows.append(int(fields[0]))
    return tuple(dict.fromkeys(rows))


def full_coverage(payload: str) -> list[Any]:
    """A clean coverage record for every graded row the payload contains.

    Built as `RowCoverage` instances rather than dicts: `model_copy(update=...)` does not
    validate, so dicts would survive all the way into the serialized request and the only
    sign would be a warning nobody reads.
    """
    from oatutor_council.agents.schemas import RowCoverage

    return [
        RowCoverage(
            row=row,
            computed_answer="checked",
            submitted_answer="checked",
            answer_correct=True,
            answer_type_correct=True,
            requested_form_correct=True,
            domain_checked=True,
            solution_count_checked=True,
            units_checked=True,
            choices_checked=True,
        )
        for row in graded_rows_in(payload)
    ]


def with_coverage(response: Any, payload: str) -> Any:
    """Fill a scan response's coverage from the block it was sent, if it takes one.

    Left alone when the response already carries coverage, so a test that is *about*
    coverage -- a short record, a duplicated row -- keeps saying what it meant to say.
    """
    fields = getattr(type(response), "model_fields", {})
    if "coverage" not in fields or getattr(response, "coverage", None):
        return response
    return response.model_copy(update={"coverage": full_coverage(payload)})


def compliant(reply: Any) -> Any:
    """Wrap a scripted reply so its scan responses carry the coverage a real one would.

    Applied at the double, not inside `ScriptedLLMClient`: the production mock must not
    fabricate a field the model is required to produce, or the one test that checks the
    requirement would be the only place the requirement existed.
    """

    def wrapped(request: Any) -> Any:
        value = reply(request) if callable(reply) else reply
        return with_coverage(value, request.user_payload)

    return wrapped
