"""Final gate tests.

The separation between integrity and content is the thing under test. Integrity failing
means the system misbehaved and nothing excuses it; content findings are routed to a
reviewer and must not, on their own, turn a completed job into a failure.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import problem, step
from oatutor_council.models import CellEdit, ColumnKey, SourcePath
from oatutor_council.validation.final_gate import GateCode, run_final_gate
from oatutor_council.workbook.writer import apply_edits, create_working_copy


@pytest.fixture
def source(make_workbook) -> Path:
    return make_workbook(
        [
            problem("angles1", title="Convert", oer_src="src", license="CC-BY"),
            step("angles1", answer="1/2", answer_type="numeric"),
            step("angles1", answer="3", answer_type="numeric"),
        ]
    )


@pytest.fixture
def copy(source, tmp_path):
    return create_working_copy(SourcePath(str(source)), tmp_path / "job")


def gate(copy, changes):
    return run_final_gate(
        source=copy.source,
        source_sha256=copy.source_sha256,
        output=copy.path,
        changes=changes,
    )


def test_an_untouched_copy_passes(copy):
    result = gate(copy, [])
    assert result.passed
    assert result.unexplained == ()


def test_an_applied_and_recorded_edit_passes(copy):
    changes = apply_edits(
        copy,
        [
            CellEdit(
                row=3, column=5, column_key=ColumnKey.ANSWER, before="1/2", after="1/3"
            )
        ],
    )
    result = gate(copy, changes)
    assert result.passed, [f.message for f in result.integrity_findings]


def test_an_edit_nobody_authorised_fails(copy):
    """The case the gate exists for. An edit is applied and recorded, and separately
    something else is changed behind the ledger's back."""
    changes = apply_edits(copy, [CellEdit(row=3, column=5, before="1/2", after="1/3")])

    workbook = load_workbook(copy.path)
    workbook.active.cell(row=4, column=5, value="tampered")
    workbook.save(copy.path)
    workbook.close()

    result = gate(copy, changes)
    assert not result.passed
    assert GateCode.UNEXPLAINED_DIFFERENCE in {
        f.code for f in result.integrity_findings
    }


def test_a_recorded_change_that_never_landed_fails(copy):
    """The opposite question, and the one a difference-only reconciliation misses: the
    ledger claims an edit the output does not contain. Reporting that job as successful
    would hand the curator a report describing a workbook that does not exist."""
    from datetime import datetime, timezone

    from oatutor_council.models import ChangeRecord

    phantom = ChangeRecord(
        change_id="c1",
        issue_id="i1",
        patch_id="p1",
        block_id="block-0000",
        row=3,
        column=5,
        column_key=ColumnKey.ANSWER,
        before="1/2",
        after="1/3",
        applied_at=datetime.now(timezone.utc),
    )
    result = gate(copy, [phantom])
    assert not result.passed
    assert [f.code for f in result.integrity_findings] == [
        GateCode.RECORDED_CHANGE_NOT_PRESENT
    ]
    assert result.changes_not_present == (phantom,)


def test_a_modified_source_fails_before_anything_else_is_reported(copy, source):
    """If the original moved, every comparison below it is against the wrong baseline.
    Reporting those results would be worse than reporting nothing."""
    source.chmod(0o644)
    source.write_bytes(source.read_bytes() + b"tampered")
    result = gate(copy, [])
    assert not result.passed
    assert [f.code for f in result.integrity_findings] == [GateCode.SOURCE_MODIFIED]
    assert result.content_findings == ()


def test_an_unreadable_output_fails(copy):
    copy.path.write_bytes(b"not a workbook at all")
    result = gate(copy, [])
    assert not result.passed
    assert [f.code for f in result.integrity_findings] == [GateCode.OUTPUT_UNREADABLE]


def test_content_findings_do_not_fail_the_integrity_gate(make_workbook, tmp_path):
    """A workbook that arrived with warnings does not become a failure because it still
    has some. Content is routed to a reviewer; integrity is not negotiable."""
    flawed = make_workbook(
        [problem("a1"), step("a1", answer="x^2", answer_type="algebra")]
    )
    copy = create_working_copy(SourcePath(str(flawed)), tmp_path / "job2")
    result = gate(copy, [])
    assert result.passed
    assert result.content_findings
    assert "content finding" in result.summary()
