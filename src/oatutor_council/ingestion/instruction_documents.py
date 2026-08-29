"""Extract text from the curator's optional instruction document.

The document *seeds* the Initial Auditor. It never replaces autonomous inspection: with
no document at all the auditor still examines every problem block, and a claim the
document makes that turns out to be false is recorded as refuted rather than quietly
dropped.

**The rule this module exists to enforce:** a document that could not be read is never
reported as a document containing no instructions. Those two outcomes look identical
downstream -- an empty seed list either way -- and confusing them means a curator who
uploaded a scanned PDF is told their workbook is fine. So extraction either produces text
or raises `UnsupportedDocumentError`; there is no path that returns nothing quietly.

What this module does *not* do is decide what the text means. Splitting prose into issue
claims is semantic judgment and belongs to the auditor. Here the text is only located,
bounded, and labelled with where it came from.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

#: Enough characters to hold any realistic errata note, and small enough that a
#: pathological upload cannot dominate a prompt. Truncation is reported, never silent.
MAX_DOCUMENT_CHARACTERS = 200_000

#: A segment longer than this is split. Long segments defeat provenance -- "somewhere in
#: this 40,000-character block" tells a curator nothing about where a claim came from.
MAX_SEGMENT_CHARACTERS = 4_000

#: Below this, a PDF's text layer is noise rather than content: scanners commonly emit a
#: handful of stray glyphs from page furniture even with no real text.
MIN_PDF_TEXT_CHARACTERS = 20

_WORD = re.compile(r"[A-Za-z]{2,}")


class DocumentFormat(StrEnum):
    MARKDOWN = "markdown"
    TEXT = "text"
    DOCX = "docx"
    PDF = "pdf"


SUPPORTED_EXTENSIONS: dict[str, DocumentFormat] = {
    ".md": DocumentFormat.MARKDOWN,
    ".markdown": DocumentFormat.MARKDOWN,
    ".txt": DocumentFormat.TEXT,
    ".docx": DocumentFormat.DOCX,
    ".pdf": DocumentFormat.PDF,
}


class UnsupportedDocumentError(Exception):
    """The document could not be read, and the curator needs to know why.

    Carries a message written for a person rather than a log: the API returns it
    verbatim, so it has to say what was wrong *and* what to do instead.
    """

    def __init__(self, message: str, *, filename: str = "") -> None:
        super().__init__(message)
        self.user_message = message
        self.filename = filename


class SegmentPurpose(StrEnum):
    """What a passage of the curator's document is *for*.

    The distinction the whole of this module turns on: "Problem 3 has the wrong answer"
    is a **hypothesis** about one block, to be checked and confirmed or refuted, while
    "Steps must not carry dependencies" is a **governing rule** that applies to every
    block and is not something any block can refute. Treating the second as the first is
    why a policy statement was being sent to thirty problems and marked refuted by
    twenty-nine of them.
    """

    #: Authoritative instruction. Goes to the Initial Auditor, Writer and both reviewers.
    RULES = "rules"
    #: A suspected defect. Goes to the Initial Auditor as a claim to verify.
    ERRATA = "errata"
    #: Context that is neither. Shown only to the auditor, never treated as a claim or
    #: policy and never passed to the Writer or reviewers.
    NOTES = "notes"


#: Phrasing that makes a passage a general instruction rather than a report about one
#: problem. Modal obligation ("must", "should never") is the strongest single signal: a
#: curator describing a specific defect says what *is* wrong, not what *must* be true.
_RULE_MARKERS = (
    "must ", "must not", "should ", "should not", "never ", "always ",
    "are required", "is required", "do not ", "don't ", "ensure ", "every ",
    "all rows", "all answers", "all problems", "convention", "policy",
)

#: A reference to a specific location. The presence of one is what makes a passage a
#: claim about a particular block rather than a statement about all of them.
_PROBLEM_REFERENCE = re.compile(
    r"\b(?:problem|question|row|line|item)\s*#?\s*([A-Za-z]*\d+)\b", re.IGNORECASE
)

#: A bare identifier that looks like a Problem Name, e.g. `angles12`, `Unitcirc3`.
_NAME_REFERENCE = re.compile(r"\b([A-Za-z]{3,}\d+)\b")

#: Language that reports a defect rather than describing one location or stating a rule.
#: Needed because a curator writing "something is wrong with one of the answers" has
#: reported a real problem in words that name neither a place nor an obligation -- and
#: filing that as background notes means nobody ever checks it.
_DEFECT_MARKERS = (
    "wrong", "incorrect", "missing", "error", "typo", "mistake", "broken",
    "does not match", "doesn't match", "no longer", "fails", "bad ",
)


def classify_segment(text: str) -> SegmentPurpose:
    """Decide what one passage is, from its own wording.

    Deliberately deterministic, and applied in a fixed order: a passage naming a specific
    problem or row is errata; a passage stating an obligation is a rule; a passage
    reporting a defect without saying where is errata too; anything left is notes, which
    are recorded and shown but never become claims a block can refute.

    Where both signals appear -- "Problem 3 must have an answer" -- the *reference* wins.
    A statement about one named problem is checkable against that problem, and treating
    it as universal policy would apply it to twenty-nine blocks it was never about.
    """
    lowered = text.lower()
    if _PROBLEM_REFERENCE.search(text) or _NAME_REFERENCE.search(text):
        return SegmentPurpose.ERRATA
    if any(marker in lowered for marker in _RULE_MARKERS):
        return SegmentPurpose.RULES
    if any(marker in lowered for marker in _DEFECT_MARKERS):
        # A defect report that named no location. It goes to every block, which is
        # wasteful and honest: the curator said something is wrong, and filing that as
        # background notes means nobody ever looks.
        return SegmentPurpose.ERRATA
    return SegmentPurpose.NOTES


def referenced_locations(text: str) -> tuple[frozenset[str], frozenset[int]]:
    """Problem names and row numbers a passage points at.

    Used to send a claim only to the blocks it could possibly be about. A claim that
    names nothing is returned empty and goes to every block, which is the honest fallback
    -- there is no way to target what the curator did not identify.
    """
    names: set[str] = set()
    rows: set[int] = set()

    for match in _PROBLEM_REFERENCE.finditer(text):
        token = match.group(1)
        if token.isdigit():
            rows.add(int(token))
        else:
            names.add(token.casefold())
    for match in _NAME_REFERENCE.finditer(text):
        names.add(match.group(1).casefold())
    return frozenset(names), frozenset(rows)


@dataclass(frozen=True)
class DocumentSegment:
    """One passage of the document, with enough provenance to cite it."""

    index: int
    text: str
    provenance: str
    purpose: SegmentPurpose = SegmentPurpose.ERRATA


@dataclass(frozen=True)
class InstructionDocument:
    filename: str
    format: DocumentFormat
    segments: tuple[DocumentSegment, ...]
    truncated: bool = False

    @property
    def character_count(self) -> int:
        return sum(len(segment.text) for segment in self.segments)

    def render(self) -> str:
        return "\n\n".join(
            f"[{segment.provenance}] {segment.text}" for segment in self.segments
        )


# --------------------------------------------------------------------------------------
# Format-specific extraction
# --------------------------------------------------------------------------------------


def _read_plain_text(path: Path) -> list[tuple[str, str]]:
    """Return `(text, provenance)` pairs from a text or Markdown file.

    Split on blank lines, because that is what separates one instruction from the next in
    every plain-text errata note. The line number is kept so a claim can be cited back to
    the document a curator actually wrote.
    """
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            content = path.read_text(encoding="latin-1")
        except Exception as error:
            raise UnsupportedDocumentError(
                f"{path.name} is not readable as text. Please supply a UTF-8 encoded "
                "TXT, Markdown, DOCX, or text-based PDF file.",
                filename=path.name,
            ) from error

    passages: list[tuple[str, str]] = []
    line_number = 1
    for chunk in re.split(r"\n\s*\n", content):
        stripped = chunk.strip()
        if stripped:
            passages.append((stripped, f"line {line_number}"))
        line_number += chunk.count("\n") + 2
    return passages


def _read_docx(path: Path) -> list[tuple[str, str]]:
    """Paragraphs and table cells.

    Table cells are included deliberately: errata notes are frequently written as a table
    of problem-name against correction, and a paragraphs-only reader would report such a
    document as empty -- the exact silent failure this module forbids.
    """
    try:
        import docx
    except ImportError as error:  # pragma: no cover - dependency is declared
        raise UnsupportedDocumentError(
            "DOCX support is unavailable in this deployment.", filename=path.name
        ) from error

    try:
        document = docx.Document(str(path))
    except Exception as error:
        raise UnsupportedDocumentError(
            f"{path.name} could not be opened as a Word document. If it is an older "
            ".doc file, please save it as .docx and upload it again.",
            filename=path.name,
        ) from error

    passages: list[tuple[str, str]] = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            passages.append((text, f"paragraph {number}"))

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                passages.append(
                    (" | ".join(cells), f"table {table_number} row {row_number}")
                )
    return passages


def _read_pdf(path: Path) -> list[tuple[str, str]]:
    """Extract the **embedded text layer** only.

    There is no OCR here and there will not be. A PDF with no text layer -- a scan, a
    screenshot, an exported image -- yields nothing, and that must surface as a clear
    error naming the supported alternatives. Returning an empty instruction set would
    tell the curator their document contained no issues, which is a different and far
    more damaging statement than "this file cannot be read".
    """
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - dependency is declared
        raise UnsupportedDocumentError(
            "PDF support is unavailable in this deployment.", filename=path.name
        ) from error

    try:
        reader = PdfReader(str(path))
        pages = list(reader.pages)
    except Exception as error:
        raise UnsupportedDocumentError(
            f"{path.name} could not be opened as a PDF. Please check the file is not "
            "corrupted or password protected.",
            filename=path.name,
        ) from error

    if not pages:
        raise UnsupportedDocumentError(
            f"{path.name} contains no pages.", filename=path.name
        )

    passages: list[tuple[str, str]] = []
    total = 0
    for number, page in enumerate(pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        stripped = text.strip()
        total += len(stripped)
        if stripped:
            passages.append((stripped, f"page {number}"))

    if total < MIN_PDF_TEXT_CHARACTERS or not _WORD.search(" ".join(t for t, _ in passages)):
        raise UnsupportedDocumentError(
            f"{path.name} has no readable text layer. Image-based and scanned PDFs are "
            "not supported, because this system does not perform OCR. Please supply a "
            "text-based PDF, or a DOCX, TXT, or Markdown file instead.",
            filename=path.name,
        )
    return passages


_READERS = {
    DocumentFormat.MARKDOWN: _read_plain_text,
    DocumentFormat.TEXT: _read_plain_text,
    DocumentFormat.DOCX: _read_docx,
    DocumentFormat.PDF: _read_pdf,
}


# --------------------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------------------


def detect_format(filename: str) -> DocumentFormat:
    extension = Path(filename).suffix.lower()
    fmt = SUPPORTED_EXTENSIONS.get(extension)
    if fmt is None:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedDocumentError(
            f"{filename} is not a supported instruction document. Supported formats "
            f"are: {supported}.",
            filename=filename,
        )
    return fmt


def read_instruction_document(
    path: Path,
    *,
    display_name: str | None = None,
    declared_purpose: SegmentPurpose | None = None,
) -> InstructionDocument:
    """Read one instruction document, or explain clearly why it cannot be read.

    `declared_purpose` is the curator saying what the whole document is. It overrides the
    per-segment classification, because someone who uploads a file and labels it a rules
    document knows something the wording does not always reveal.
    """
    name = display_name or path.name

    if not path.is_file():
        raise UnsupportedDocumentError(
            f"{name} could not be found.", filename=name
        )
    if path.stat().st_size == 0:
        raise UnsupportedDocumentError(
            f"{name} is empty. An empty file cannot be distinguished from a document "
            "that failed to upload, so it is rejected rather than treated as containing "
            "no instructions.",
            filename=name,
        )

    fmt = detect_format(name)
    passages = _READERS[fmt](path)

    if not passages:
        raise UnsupportedDocumentError(
            f"{name} contains no readable text. If the document is a scan or a set of "
            "screenshots, this system cannot read it -- there is no OCR. Please supply "
            "a text-based PDF, DOCX, TXT, or Markdown file.",
            filename=name,
        )

    segments, truncated = _bound(passages, declared_purpose)
    return InstructionDocument(
        filename=name, format=fmt, segments=segments, truncated=truncated
    )


def _bound(
    passages: list[tuple[str, str]],
    declared: SegmentPurpose | None = None,
) -> tuple[tuple[DocumentSegment, ...], bool]:
    """Split oversized passages and stop at the document budget.

    Truncation is returned as a flag rather than hidden, so a report can say that the
    document was cut off. Silently dropping the tail of an errata note would lose real
    instructions with no trace.
    """
    segments: list[DocumentSegment] = []
    used = 0
    truncated = False

    for text, provenance in passages:
        for part, label in _split(text, provenance):
            if used + len(part) > MAX_DOCUMENT_CHARACTERS:
                truncated = True
                return tuple(segments), truncated
            segments.append(
                DocumentSegment(
                    len(segments),
                    part,
                    label,
                    declared or classify_segment(part),
                )
            )
            used += len(part)
    return tuple(segments), truncated


def _split(text: str, provenance: str):
    if len(text) <= MAX_SEGMENT_CHARACTERS:
        yield text, provenance
        return
    for offset in range(0, len(text), MAX_SEGMENT_CHARACTERS):
        part = text[offset : offset + MAX_SEGMENT_CHARACTERS]
        yield part, f"{provenance} (continued)" if offset else provenance
