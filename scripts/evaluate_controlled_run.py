#!/usr/bin/env python3
"""Compare a corrected controlled workbook with its source and hidden answer key.

This script makes pilot quality measurable. It never calls a model and never writes a
workbook. Keys may require one exact value, accept several valid representations, compare
mathematical equivalence, or assert a narrow invariant such as multiple-choice exactness.
That distinction matters: a single preferred distractor is not the only correct repair,
and an evaluator that says otherwise trains the project toward its own synthetic answers.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from oatutor_council.validation.mathematics import (  # noqa: E402
    MathVerdict,
    answers_equivalent,
    equations_equivalent,
)


def rendered(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def keyed_items(key: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    automated: list[dict[str, Any]] = []
    manual: list[dict[str, Any]] = []
    if "defectGroups" in key:
        for group in key["defectGroups"]:
            for correction in group.get("corrections", []):
                item = {
                    **correction,
                    "problemName": group["problemName"],
                    "kind": group["kind"],
                }
                if _has_automated_check(item):
                    automated.append(item)
                else:
                    manual.append(item)
        return automated, manual

    for defect in key.get("plantedDefects", []):
        if not defect.get("cell"):
            manual.append(defect)
            continue
        candidate = defect.get("suggested", defect.get("expected"))
        if not isinstance(candidate, (str, int, float)):
            manual.append(defect)
            continue
        # Sentences describe a property, not one character-exact acceptable value.
        if isinstance(candidate, str) and (" " in candidate.strip() or candidate.endswith(".")):
            manual.append(defect)
            continue
        automated.append({**defect, "expected": candidate})
    return automated, manual


def supplementary_items(key: dict[str, Any]) -> list[dict[str, Any]]:
    """Known source defects are allowed and reported, but never inflate the planted score."""
    items: list[dict[str, Any]] = []
    for group in key.get("knownSourceDefects", []):
        for correction in group.get("corrections", []):
            items.append(
                {
                    **correction,
                    "problemName": group["problemName"],
                    "kind": group["kind"],
                }
            )
    return items


def allowed_related_cells(key: dict[str, Any]) -> set[str]:
    """Cells a complete planted repair may change without becoming separate score items."""
    return {
        str(cell)
        for group in key.get("defectGroups", [])
        for cell in group.get("allowedRelatedCells", [])
        if cell
    }


def _has_automated_check(item: dict[str, Any]) -> bool:
    """Whether a key supplies a machine-decidable success condition.

    A prose hint is not machine-decidable merely because its author wrote one preferred
    sentence in ``expected``. That mistake made two substantively correct repairs fail a
    sealed score character-for-character. Predicates, accepted sets and mathematical
    equivalence are explicit contracts. An exact prose string must likewise opt in with
    ``comparison: exact``; otherwise title/body prose stays in the manual queue.
    """
    if not item.get("cell"):
        return False
    if any(key in item for key in ("accepted", "mathEquivalentTo", "predicate")):
        return True
    if "expected" not in item:
        return False
    expected = item.get("expected")
    if item.get("comparison") == "exact":
        return isinstance(expected, (str, int, float))
    if not isinstance(expected, (str, int, float)):
        return False
    coordinate = str(item["cell"]).upper()
    column = "".join(character for character in coordinate if character.isalpha())
    kind = str(item.get("kind") or "").casefold()
    instructional_prose = column in {"C", "D"} and any(
        marker in kind for marker in ("hint", "instruction", "wording", "explanation")
    )
    return not (instructional_prose and isinstance(expected, str))


def evaluate_item(item: dict[str, Any], sheet) -> tuple[bool, str, str]:
    """Return ``(passed, actual, explanation)`` for one key item."""
    coordinate = item["cell"]
    actual = rendered(sheet[coordinate].value)

    if "accepted" in item:
        accepted = [rendered(value) for value in item["accepted"]]
        return actual in accepted, actual, f"accepted values: {accepted!r}"
    if "mathEquivalentTo" in item:
        expected = rendered(item["mathEquivalentTo"])
        passed = answers_equivalent(actual, expected) is MathVerdict.EQUIVALENT
        return passed, actual, f"mathematically equivalent to {expected!r}"
    if "predicate" in item:
        passed, explanation = _predicate(item["predicate"], coordinate, actual, sheet)
        return passed, actual, explanation

    expected = rendered(item.get("expected"))
    return actual == expected, actual, f"exactly {expected!r}"


def _predicate(name: str, coordinate: str, actual: str, sheet) -> tuple[bool, str]:
    if name == "blank":
        return actual == "", "cell is blank"
    if name == "nonempty":
        return bool(actual), "cell is non-empty"
    if name == "trimmed":
        return actual == actual.strip(), "cell has no boundary whitespace"
    if name == "ascii_only":
        return all(ord(character) < 128 for character in actual), "cell is ASCII-only"
    if name == "mc_exact_answer_once":
        row = sheet[coordinate].row
        answer = rendered(sheet.cell(row=row, column=5).value)
        choices = actual.split("|") if actual else []
        exact_count = choices.count(answer)
        equivalent_distractors = [
            choice
            for choice in choices
            if choice != answer
            and equations_equivalent(answer, choice) is MathVerdict.EQUIVALENT
        ]
        passed = 2 <= len(choices) <= 5 and exact_count == 1 and not equivalent_distractors
        return (
            passed,
            "2-5 choices, exact Answer appears once, and no distractor is equivalent",
        )
    raise ValueError(f"unknown evaluation predicate {name!r}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("corrected", type=Path)
    parser.add_argument("key", type=Path)
    parser.add_argument(
        "--summary",
        action="store_true",
        help="print a concise score instead of the complete JSON result",
    )
    args = parser.parse_args()

    key = json.loads(args.key.read_text(encoding="utf-8"))
    source_book = load_workbook(args.source, data_only=False)
    corrected_book = load_workbook(args.corrected, data_only=False)
    try:
        source = source_book.active
        corrected = corrected_book.active
        automated, manual = keyed_items(key)
        supplementary = supplementary_items(key)
        target_cells = {
            item["cell"]
            for item in (*automated, *manual, *supplementary)
            if item.get("cell")
        } | allowed_related_cells(key)

        passed = []
        failed = []
        for item in automated:
            item_passed, actual, check = evaluate_item(item, corrected)
            row = {**item, "actual": actual, "check": check}
            (passed if item_passed else failed).append(row)

        supplementary_results = []
        for item in supplementary:
            if _has_automated_check(item):
                item_passed, actual, check = evaluate_item(item, corrected)
                supplementary_results.append(
                    {**item, "passed": item_passed, "actual": actual, "check": check}
                )
            else:
                supplementary_results.append({**item, "passed": None})

        unexpected = []
        max_row = max(source.max_row, corrected.max_row)
        max_column = max(source.max_column, corrected.max_column)
        for row in range(1, max_row + 1):
            for column in range(1, max_column + 1):
                old = rendered(source.cell(row=row, column=column).value)
                new = rendered(corrected.cell(row=row, column=column).value)
                coordinate = corrected.cell(row=row, column=column).coordinate
                if old != new and coordinate not in target_cells:
                    unexpected.append({"cell": coordinate, "before": old, "after": new})

        clean_names = {
            str(item.get("problemName", "")).casefold()
            for item in key.get("cleanControls", [])
            if item.get("problemName")
        }
        changed_clean_controls = sorted(
            {
                rendered(source.cell(row=corrected[item["cell"]].row, column=1).value)
                for item in unexpected
                if rendered(
                    source.cell(row=corrected[item["cell"]].row, column=1).value
                ).casefold()
                in clean_names
            }
        )

        result = {
            "source": str(args.source),
            "corrected": str(args.corrected),
            "key": str(args.key),
            "automated_passed": len(passed),
            "automated_total": len(automated),
            "automated_failed": failed,
            # Backward-compatible field names for saved reports and older tooling. They
            # now count all automated checks, not only character-exact values.
            "exact_passed": len(passed),
            "exact_total": len(automated),
            "exact_failed": failed,
            "manual_expectations": manual,
            "known_source_defects": supplementary_results,
            "unexpected_changed_cells": unexpected,
            "changed_clean_controls": changed_clean_controls,
        }
        if args.summary:
            failed_cells = ", ".join(item["cell"] for item in failed) or "none"
            unexpected_cells = ", ".join(item["cell"] for item in unexpected) or "none"
            print(f"automated: {len(passed)}/{len(automated)}")
            print(f"failed automated cells: {failed_cells}")
            print(f"manual expectations: {len(manual)}")
            print(f"known source defects: {len(supplementary_results)} (not scored)")
            print(f"unexpected changed cells: {unexpected_cells}")
            print(
                "changed clean controls: "
                + (", ".join(changed_clean_controls) or "none")
            )
        else:
            print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0 if not failed and not unexpected else 1
    finally:
        source_book.close()
        corrected_book.close()


if __name__ == "__main__":
    raise SystemExit(main())
