"""Grade a corrected workbook against a hidden before/after reference pair.

This is intentionally independent of the council database. It answers whether the file
handed back reached the human reference, including misses and unexpected edits, rather
than whether the pipeline successfully closed the issues it happened to open.
"""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class CellGrade:
    coordinate: str
    source: Any
    expected: Any
    actual: Any
    outcome: str


def _values(path: Path) -> tuple[str, dict[tuple[int, int], Any]]:
    # Normal mode deliberately: some standards-compliant writers omit the optional
    # worksheet-dimension hint, which makes openpyxl's read-only max_row/max_column None.
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook.active
    values = {
        (row, column): sheet.cell(row=row, column=column).value
        for row in range(1, sheet.max_row + 1)
        for column in range(1, sheet.max_column + 1)
    }
    title = sheet.title
    workbook.close()
    return title, values


def grade(source: Path, expected: Path, actual: Path) -> tuple[CellGrade, ...]:
    source_sheet, before = _values(source)
    expected_sheet, reference = _values(expected)
    actual_sheet, result = _values(actual)
    if len({source_sheet, expected_sheet, actual_sheet}) != 1:
        raise ValueError(
            f"active sheets differ: {source_sheet!r}, {expected_sheet!r}, {actual_sheet!r}"
        )

    grades: list[CellGrade] = []
    for row, column in sorted(set(before) | set(reference) | set(result)):
        old = before.get((row, column))
        wanted = reference.get((row, column))
        got = result.get((row, column))
        if old == wanted and old == got:
            continue
        if old != wanted:
            if got == wanted:
                outcome = "matched_reference"
            elif got == old:
                outcome = "missed"
            else:
                outcome = "different_correction"
        else:
            outcome = "unexpected_edit"
        grades.append(
            CellGrade(
                coordinate=f"{get_column_letter(column)}{row}",
                source=old,
                expected=wanted,
                actual=got,
                outcome=outcome,
            )
        )
    return tuple(grades)


def summary(grades: tuple[CellGrade, ...]) -> dict[str, Any]:
    expected = [grade for grade in grades if grade.source != grade.expected]
    actual = [grade for grade in grades if grade.source != grade.actual]
    matched = [grade for grade in expected if grade.outcome == "matched_reference"]
    return {
        "expected_changed_cells": len(expected),
        "actual_changed_cells": len(actual),
        "matched_reference": len(matched),
        "missed": sum(grade.outcome == "missed" for grade in grades),
        "different_correction": sum(
            grade.outcome == "different_correction" for grade in grades
        ),
        "unexpected_edit": sum(grade.outcome == "unexpected_edit" for grade in grades),
        "target_recall": len(matched) / len(expected) if expected else 1.0,
        "edit_precision": len(matched) / len(actual) if actual else 1.0,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("expected", type=Path)
    parser.add_argument("actual", type=Path)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    grades = grade(args.source, args.expected, args.actual)
    report = {"summary": summary(grades), "cells": [asdict(item) for item in grades]}
    if args.json:
        print(json.dumps(report, indent=2, default=str))
    else:
        print(json.dumps(report["summary"], indent=2))
        for item in grades:
            if item.outcome != "matched_reference":
                print(
                    f"{item.coordinate}: {item.outcome}\n"
                    f"  source={item.source!r}\n  expected={item.expected!r}\n"
                    f"  actual={item.actual!r}"
                )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
