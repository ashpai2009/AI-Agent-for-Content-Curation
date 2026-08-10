#!/usr/bin/env python3
"""Read-only reconnaissance across the real OATutor workbook corpus.

Answers the questions the reader and rule engine must be built against: where the
header actually sits, how blocks are really delimited, which conventions each
workbook uses, and which defects are genuinely present.

This script NEVER writes to a source workbook. Every file is hashed before and
after inspection and the run fails loudly if any hash moves. Output goes to
`recon/`, which is gitignored.

    python scripts/recon_workbooks.py [--corpus ~/Documents/OATutor] [--out recon]
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

# The column contract we expect. Index is 1-based to match openpyxl.
EXPECTED_HEADERS = [
    "Problem Name", "Row Type", "Title", "Body Text", "Answer", "answerType",
    "HintID", "Dependency", "mcChoices", "Images (space delimited)", "Parent",
    "OER src", "openstax KC", "KC", "Taxonomy", "License", None, None,
    "Validator Check", "Time Last Checked",
]
COL = {name: i + 1 for i, name in enumerate(
    ["problem_name", "row_type", "title", "body_text", "answer", "answer_type",
     "hint_id", "dependency", "mc_choices", "images", "parent", "oer_src",
     "openstax_kc", "kc", "taxonomy", "license", "q", "r", "validator", "checked"]
)}

ID_RE = re.compile(r"^([a-zA-Z]+)(\d+)$")
STEM_RE = re.compile(r"^(.*?)(\d+)$")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def is_blank_row(ws, row: int, last_col: int) -> bool:
    return all(ws.cell(row=row, column=c).value in (None, "") for c in range(1, last_col + 1))


def find_header_row(ws, last_col: int, scan_limit: int = 15) -> int | None:
    """The header is wherever 'Problem Name' actually appears, not row 1 by fiat."""
    for row in range(1, min(scan_limit, ws.max_row) + 1):
        for col in range(1, last_col + 1):
            if cell_text(ws.cell(row=row, column=col).value).strip() == "Problem Name":
                return row
    return None


def segment_blocks(ws, first_data_row: int, last_row: int) -> tuple[list[dict], list[dict]]:
    """Segment by `Row Type == problem`, then validate Problem Name agreement.

    This is the rule the production reader uses. Running it here tells us whether
    real workbooks actually satisfy it, and surfaces the disagreements rather than
    silently resolving them.
    """
    blocks: list[dict] = []
    anomalies: list[dict] = []

    problem_rows = [
        r for r in range(first_data_row, last_row + 1)
        if cell_text(ws.cell(row=r, column=COL["row_type"]).value).strip().lower() == "problem"
    ]

    # Rows carrying content before the first problem row are orphans.
    for r in range(first_data_row, problem_rows[0] if problem_rows else last_row + 1):
        if not is_blank_row(ws, r, 20):
            anomalies.append({"code": "ORPHAN_ROW_BEFORE_FIRST_PROBLEM", "row": r})

    for i, start in enumerate(problem_rows):
        end = (problem_rows[i + 1] - 1) if i + 1 < len(problem_rows) else last_row
        # Trim trailing blank separator rows off the block, but do not let a blank
        # row in the middle truncate it.
        while end > start and is_blank_row(ws, end, 20):
            end -= 1

        expected = cell_text(ws.cell(row=start, column=COL["problem_name"]).value).strip()
        interior_blanks, name_mismatches, missing_names = [], [], []
        for r in range(start, end + 1):
            if is_blank_row(ws, r, 20):
                interior_blanks.append(r)
                continue
            got = cell_text(ws.cell(row=r, column=COL["problem_name"]).value).strip()
            if not got:
                # A blank Problem Name is ambiguous: the value may simply be absent,
                # or the whole row may be shifted right so the name landed in Title.
                # Distinguishing them matters because the repairs are different.
                shifted = cell_text(ws.cell(row=r, column=COL["title"]).value).strip()
                if (not cell_text(ws.cell(row=r, column=COL["row_type"]).value).strip()
                        and shifted == expected):
                    anomalies.append({"code": "ROW_SHIFT_RIGHT", "row": r, "block": expected,
                                      "detail": "Problem Name found in Title column"})
                else:
                    missing_names.append(r)
            elif got != expected:
                name_mismatches.append({"row": r, "expected": expected, "got": got})

        blocks.append({
            "problem_name": expected,
            "start_row": start,
            "end_row": end,
            "rows": end - start + 1,
            "interior_blank_rows": interior_blanks,
        })
        for m in name_mismatches:
            anomalies.append({"code": "PROBLEM_NAME_MISMATCH_IN_BLOCK", **m})
        for r in missing_names:
            anomalies.append({"code": "MISSING_PROBLEM_NAME", "row": r, "block": expected})
        if interior_blanks:
            anomalies.append({
                "code": "INTERIOR_BLANK_ROW", "block": expected, "rows": interior_blanks,
            })

    return blocks, anomalies


def detect_dependency_convention(ws, blocks: list[dict]) -> dict:
    """reset-per-step restarts hint ids at h1 on each step; continuous keeps climbing."""
    reset, continuous, undecided = 0, 0, 0
    for b in blocks:
        seq_by_step: list[list[int]] = []
        current: list[int] = []
        for r in range(b["start_row"], b["end_row"] + 1):
            rt = cell_text(ws.cell(row=r, column=COL["row_type"]).value).strip().lower()
            if rt == "step":
                if current:
                    seq_by_step.append(current)
                current = []
            elif rt == "hint":
                m = ID_RE.match(cell_text(ws.cell(row=r, column=COL["hint_id"]).value).strip())
                if m:
                    current.append(int(m.group(2)))
        if current:
            seq_by_step.append(current)

        multi = [s for s in seq_by_step if s]
        if len(multi) < 2:
            undecided += 1
        elif all(s and s[0] == 1 for s in multi):
            reset += 1
        elif all(
            multi[i][-1] < multi[i + 1][0] for i in range(len(multi) - 1) if multi[i] and multi[i + 1]
        ):
            continuous += 1
        else:
            undecided += 1

    verdict = "reset" if reset > continuous else "continuous" if continuous > reset else "undecided"
    return {"verdict": verdict, "reset_blocks": reset,
            "continuous_blocks": continuous, "undecided_blocks": undecided}


def inspect(path: Path) -> dict:
    before = sha256(path)
    wb = openpyxl.load_workbook(path, data_only=False)  # keep formulas; never save
    try:
        ws = wb[wb.sheetnames[0]]
        last_col, last_row = 20, ws.max_row

        header_row = find_header_row(ws, last_col)
        headers = ([cell_text(ws.cell(row=header_row, column=c).value) or None
                    for c in range(1, last_col + 1)] if header_row else [])

        first_data_row = None
        if header_row:
            r = header_row + 1
            while r <= last_row and is_blank_row(ws, r, last_col):
                r += 1
            first_data_row = r if r <= last_row else None

        report: dict[str, Any] = {
            "file": path.name,
            "sha256": before,
            "sheet_names": wb.sheetnames,
            "sheet_count": len(wb.sheetnames),
            "max_row": last_row,
            "header_row": header_row,
            "headers_match_contract": headers == EXPECTED_HEADERS if headers else None,
            "headers": headers,
            "first_data_row": first_data_row,
            "blank_rows_between_header_and_data": (
                (first_data_row - header_row - 1) if header_row and first_data_row else None
            ),
        }
        if not first_data_row:
            report["error"] = "no data rows found"
            return report

        blocks, anomalies = segment_blocks(ws, first_data_row, last_row)
        report["block_count"] = len(blocks)
        report["blocks_sample"] = blocks[:3]

        # --- conventions -----------------------------------------------------
        row_types = Counter()
        answer_types = Counter()
        scaffold_prefixes = Counter()
        hint_prefixes = Counter()
        latex_cells = 0
        mc_answer_first = mc_total = 0

        for r in range(first_data_row, last_row + 1):
            rt = cell_text(ws.cell(row=r, column=COL["row_type"]).value).strip()
            if rt:
                row_types[rt.lower()] += 1
            at = cell_text(ws.cell(row=r, column=COL["answer_type"]).value).strip()
            if at:
                answer_types[at] += 1
            ident = cell_text(ws.cell(row=r, column=COL["hint_id"]).value).strip()
            m = ID_RE.match(ident)
            if m:
                if rt.lower() == "scaffold":
                    scaffold_prefixes[m.group(1)] += 1
                elif rt.lower() == "hint":
                    hint_prefixes[m.group(1)] += 1
            for c in range(1, last_col + 1):
                if "$$" in cell_text(ws.cell(row=r, column=c).value):
                    latex_cells += 1
            if at == "mc":
                choices = cell_text(ws.cell(row=r, column=COL["mc_choices"]).value).split("|")
                ans = cell_text(ws.cell(row=r, column=COL["answer"]).value)
                if choices and choices[0]:
                    mc_total += 1
                    if choices[0].strip() == ans.strip():
                        mc_answer_first += 1

        names = [b["problem_name"] for b in blocks if b["problem_name"]]
        stems = Counter(m.group(1) for n in names if (m := STEM_RE.match(n)))

        report["row_types"] = dict(row_types)
        report["answer_types"] = dict(answer_types)
        report["naming_stems"] = dict(stems)
        report["scaffold_id_prefixes"] = dict(scaffold_prefixes)
        report["hint_id_prefixes"] = dict(hint_prefixes)
        report["latex_cell_count"] = latex_cells
        report["notation"] = ("latex" if latex_cells > 50 else
                              "ascii" if latex_cells == 0 else "mixed")
        report["mc_answer_is_first_choice"] = {"count": mc_answer_first, "of": mc_total}
        report["dependency_convention"] = detect_dependency_convention(ws, blocks)

        # --- defects ---------------------------------------------------------
        for r in range(first_data_row, last_row + 1):
            for key in ("answer", "mc_choices"):
                v = ws.cell(row=r, column=COL[key]).value
                if isinstance(v, (dt.datetime, dt.date)):
                    anomalies.append({"code": "DATE_COERCION", "row": r,
                                      "column": key, "value": str(v)})
            at_val = cell_text(ws.cell(row=r, column=COL["answer_type"]).value).strip()
            if at_val and ID_RE.match(at_val) and at_val.lower()[0] in "hs":
                anomalies.append({"code": "COLUMN_SHIFT", "row": r,
                                  "answer_type_cell": at_val})
            for c in range(1, last_col + 1):
                t = cell_text(ws.cell(row=r, column=c).value)
                if "\\\\" in t:
                    anomalies.append({"code": "DOUBLE_ESCAPED_BACKSLASH", "row": r, "column": c})
                if t.count("$$") % 2 == 1:
                    anomalies.append({"code": "UNBALANCED_LATEX_DELIMITER", "row": r, "column": c})
            # A `\middle|` inside a LaTeX choice gets eaten by pipe-splitting, which
            # destroys both the delimiters and the choice boundaries. The whole cell
            # can still have an even `$$` count, so this must be checked per choice.
            choices_cell = cell_text(ws.cell(row=r, column=COL["mc_choices"]).value)
            if "$$" in choices_cell:
                broken = [i for i, part in enumerate(choices_cell.split("|"))
                          if part.count("$$") % 2 == 1]
                if broken:
                    anomalies.append({"code": "MC_LATEX_PIPE_CORRUPTION", "row": r,
                                      "broken_choice_indexes": broken})

        counts = Counter(a["code"] for a in anomalies)
        report["anomaly_counts"] = dict(counts)
        report["anomalies_sample"] = anomalies[:25]

        # --- appearance ------------------------------------------------------
        wrap_rows, heights, fonts = set(), Counter(), Counter()
        for r in range(first_data_row, min(last_row, first_data_row + 400) + 1):
            h = ws.row_dimensions[r].height
            heights[h] += 1
            for c in range(1, last_col + 1):
                cell = ws.cell(row=r, column=c)
                if cell.alignment and cell.alignment.wrap_text:
                    wrap_rows.add(r)
                if cell.font and cell.font.name:
                    fonts[(cell.font.name, cell.font.size)] += 1
        report["appearance"] = {
            "rows_with_wrap_text": len(wrap_rows),
            "row_heights": {str(k): v for k, v in heights.most_common(6)},
            "column_widths": {k: round(v.width, 2) for k, v in
                              sorted(ws.column_dimensions.items()) if v.width},
            "fonts": {f"{n} {s}": c for (n, s), c in fonts.most_common(4)},
            "freeze_panes": ws.freeze_panes,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges][:10],
            "merged_range_count": len(ws.merged_cells.ranges),
            "hidden_rows": [r for r, d in ws.row_dimensions.items() if d.hidden][:20],
            "hidden_cols": [c for c, d in ws.column_dimensions.items() if d.hidden][:20],
            "sheet_state": ws.sheet_state,
            "data_validation_count": len(getattr(ws.data_validations, "dataValidation", [])),
            "image_count": len(getattr(ws, "_images", [])),
        }
        return report
    finally:
        wb.close()
        after = sha256(path)
        if after != before:
            raise SystemExit(f"FATAL: {path.name} changed during inspection ({before} -> {after})")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--corpus", default="~/Documents/OATutor")
    ap.add_argument("--out", default="recon")
    args = ap.parse_args()

    corpus = Path(args.corpus).expanduser()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    files = sorted(p for p in corpus.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        print(f"no .xlsx found under {corpus}", file=sys.stderr)
        return 1

    reports = []
    for p in files:
        print(f"  inspecting {p.name} ...", flush=True)
        try:
            reports.append(inspect(p))
        except SystemExit:
            raise
        except Exception as exc:  # a corrupt workbook must not abort the sweep
            reports.append({"file": p.name, "error": f"{type(exc).__name__}: {exc}"})

    (out / "workbooks.json").write_text(json.dumps(reports, indent=2, default=str))
    print(f"\nwrote {out / 'workbooks.json'}  ({len(reports)} workbooks)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
