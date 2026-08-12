"""End-to-end demonstration of the council, with no credentials and no network.

Generates a synthetic workbook containing planted defects, runs the complete five-stage
pipeline against a scripted model client, and prints every artefact a curator receives.

The mathematics is invented. Nothing here comes from the real corpus, and the workbook is
written to a temporary directory that is cleaned up on exit.

Usage:  .venv/bin/python scripts/demo.py [--keep]
"""

from __future__ import annotations

import argparse
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from openpyxl import Workbook  # noqa: E402

from oatutor_council.agents.schemas import (  # noqa: E402
    AuditorResponse,
    IndependentReviewResponse,
    ReviewerResponse,
    WriterResponse,
)
from oatutor_council.config import Settings  # noqa: E402
from oatutor_council.council import CurationCouncil  # noqa: E402
from oatutor_council.llm.base import AgentRole  # noqa: E402
from oatutor_council.llm.mock import ScriptedLLMClient  # noqa: E402
from oatutor_council.models import ColumnKey, CurationJob, SourcePath  # noqa: E402
from oatutor_council.persistence import (  # noqa: E402
    Database,
    create_job,
    list_changes,
    list_events,
    describe_artifacts,
    latest_findings,
    list_issues,
    list_llm_calls,
    list_verdicts,
    rediscovery_counts,
    token_usage,
)
from oatutor_council.reporting.reports import build_reports, render_markdown  # noqa: E402
from oatutor_council.workbook.reader import read_workbook  # noqa: E402
from oatutor_council.workbook.writer import create_working_copy, sha256_of  # noqa: E402

HEADERS = [
    "Problem Name", "Row Type", "Title", "Body Text", "Answer", "answerType",
    "HintID", "Dependency", "mcChoices", "Images (space delimited)", "Parent",
    "OER src", "openstax KC", "KC", "Taxonomy", "License", None, None,
    "Validator Check", "Time Last Checked",
]


def build_demo_workbook(path: Path) -> Path:
    """A synthetic workbook with four planted defects, each of a different kind."""
    # Every row is well-formed apart from the planted defect it carries. That is the
    # point of a demonstration: a fixture with incidental faults produces a report in
    # which the planted defects are indistinguishable from the sloppiness around them.
    rows = [
        # A sound block, to show that correct problems are left alone.
        ["conv1", "problem", "Convert an angle to radians", None, None, None, None,
         None, None, None, None, "example.org/1", "Angles", "Angles", "OpenStax", "CC-BY"],
        ["conv1", "step", "Convert 30 degrees to radians.", "Use the conversion factor.",
         "pi/6", "algebra"],
        ["conv1", "hint", "How are degrees converted?", "Multiply by pi/180.",
         None, None, "h1"],

        # Defect 1: a scaffold with no answer. Deterministic rules catch this.
        ["conv2", "problem", "Evaluate a trigonometric value", None, None, None, None,
         None, None, None, None, "example.org/2", "Trig", "Trig", "OpenStax", "CC-BY"],
        ["conv2", "step", "Evaluate cos(theta) at theta=0.", "Read the value directly.",
         "1", "numeric"],
        ["conv2", "scaffold", "What is cos(0)?", "Start from the definition.",
         None, "numeric", "s1"],

        # Defect 2: Excel coerced the fraction 1/2 into a date.
        ["conv3", "problem", "State a fractional value", None, None, None, None,
         None, None, None, None, "example.org/3", "Frac", "Frac", "OpenStax", "CC-BY"],
        ["conv3", "step", "Give the value as a fraction.", "Do not use a decimal.",
         datetime(2026, 1, 2), "numeric"],

        # Defect 3: the answer matches no choice exactly, though one is equivalent.
        ["conv4", "problem", "Choose the correct value", None, None, None, None,
         None, None, None, None, "example.org/4", "MC", "MC", "OpenStax", "CC-BY"],
        ["conv4", "step", "Which equals one half?", "Choose exactly one.", "1/2", "mc",
         None, None, "0.5|1/3|1/4"],

        # Defect 4: a dependency pointing at an identifier that does not exist.
        ["conv5", "problem", "Follow the scaffolded steps", None, None, None, None,
         None, None, None, None, "example.org/5", "Dep", "Dep", "OpenStax", "CC-BY"],
        ["conv5", "step", "Work through the parts.", "Take them in order.", "4",
         "numeric"],
        ["conv5", "scaffold", "First part.", "Halve the total.", "2", "numeric", "s1",
         "s9"],
    ]

    workbook = Workbook()
    sheet = workbook.active
    for column, label in enumerate(HEADERS, start=1):
        if label:
            sheet.cell(row=1, column=column, value=label)
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=2 + offset, column=column, value=value)
    workbook.save(path)
    workbook.close()
    return path


def scripted_client(db: Database) -> ScriptedLLMClient:
    """A model that behaves the way a competent one would on this workbook."""

    def reply(request):
        if request.role is AgentRole.INITIAL_AUDITOR:
            # The deterministic rules already found the planted defects, so the auditor
            # has nothing further to add. Zero findings is a valid and common answer.
            return AuditorResponse(reasoning="the block matches the stated mathematics")

        if request.role is AgentRole.INDEPENDENT_REVIEWER:
            return IndependentReviewResponse(block_is_sound=True)

        if request.role is AgentRole.WRITER:
            return _writer_reply(db, request)

        return ReviewerResponse(
            decision="accept", rule_codes=["MC_ANSWER_NOT_IN_CHOICES"]
        )

    client = ScriptedLLMClient()
    client.default = reply
    return client


def _writer_reply(db: Database, request) -> WriterResponse:
    """Repairs keyed to the issue's own description, as a real Writer would work."""
    issue = next(
        (i for i in list_issues(db, request.job_id) if i.issue_id == request.issue_id),
        None,
    )
    description = issue.description if issue else ""
    row = issue.cells[0][0] if issue and issue.cells else 0

    if "SCAFFOLD_MISSING_ANSWER" in (issue.title if issue else ""):
        return WriterResponse(
            reasoning="the scaffold asks for cos(0), which is 1",
            derivation="cos(0) = 1 by the definition of the cosine at zero",
            edits=[{"row": row, "column": "answer", "before": "", "after": "1"}],
        )
    if "DATE_COERCION" in (issue.title if issue else ""):
        return WriterResponse(
            reasoning="Excel turned the fraction into a date on entry",
            derivation="the stored date 2026-01-02 encodes the fraction 1/2",
            edits=[
                {
                    "row": row,
                    "column": "answer",
                    "before": "2026-01-02 00:00:00",
                    "after": "1/2",
                }
            ],
        )
    if "MC_ANSWER_NOT_IN_CHOICES" in (issue.title if issue else ""):
        return WriterResponse(
            reasoning="the answer is equal to a choice but written differently",
            derivation="1/2 and 0.5 are the same value; the choice list must match "
            "the answer exactly",
            # The issue names the Answer cell; the repair belongs in mcChoices. The gate
            # requires that deviation to be stated, and then checks that the edit was
            # actually needed rather than taking the sentence on trust.
            related_edits_reason="the answer is correct as written, so the choice list "
            "is what has to change for one choice to match it exactly",
            edits=[
                {
                    "row": row,
                    "column": "mc_choices",
                    "before": "0.5|1/3|1/4",
                    "after": "1/2|1/3|1/4",
                }
            ],
        )
    if "DEPENDENCY_UNRESOLVED" in (issue.title if issue else ""):
        return WriterResponse(
            reasoning="the dependency names an identifier the block does not contain",
            derivation="the block has only s1, so the stray reference is removed",
            edits=[{"row": row, "column": "dependency", "before": "s9", "after": ""}],
        )

    return WriterResponse(
        needs_human_review=True,
        human_review_reason=f"no confident repair for: {description[:80]}",
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--keep", action="store_true", help="keep the demo directory")
    args = parser.parse_args()

    workspace = Path(tempfile.mkdtemp(prefix="oatutor-demo-"))
    try:
        source = build_demo_workbook(workspace / "demo-workbook.xlsx")
        source_hash = sha256_of(source)

        print("=" * 78)
        print("OATutor curation council — offline demonstration")
        print("=" * 78)
        print(f"\nSubmitted workbook: {source.name}")
        parsed = read_workbook(source)
        print(f"  {len(parsed.blocks)} problem block(s), notation {parsed.conventions.notation.value}")

        db = Database(workspace / "council.db")
        copy = create_working_copy(SourcePath(str(source)), workspace / "job")
        create_job(
            db,
            CurationJob(
                job_id="demo",
                source_filename=source.name,
                source_sha256=copy.source_sha256,
            ),
        )
        settings = Settings(
            claude_cli_path="not-needed-for-the-mock",
            claude_model="mock",
            claude_effort="medium",
            data_root=workspace,
            max_repair_attempts=3,
            max_validation_rounds=2,
            step_budget=500,
            llm_call_budget=300,
            interrupted_retry_budget=2,
            max_concurrent_jobs=1,
            max_upload_bytes=52_428_800,
            lease_seconds=60,
        )

        council = CurationCouncil(
            db=db, settings=settings, client=scripted_client(db), job_id="demo", copy=copy
        )

        print("\n--- pipeline ------------------------------------------------------")
        while not council.job.is_terminal:
            outcome = council.step()
            print(f"  {outcome.state.value:22s} {outcome.description[:70]}")

        job = council.job
        changes = list_changes(db, "demo")
        issues = list_issues(db, "demo")

        print("\n--- cell changes --------------------------------------------------")
        for change in sorted(changes, key=lambda c: c.row):
            key = change.column_key.value if change.column_key else change.column
            print(f"  row {change.row:>3} {key:12s} {change.before!r} -> {change.after!r}")

        print("\n--- issue ledger --------------------------------------------------")
        for issue in issues:
            print(f"  [{issue.state.value:18s}] {issue.title}")

        print("\n--- reviewer decisions --------------------------------------------")
        for verdict in list_verdicts(db, "demo"):
            print(
                f"  {verdict.reviewer_role.value:22s} attempt {verdict.attempt_no}: "
                f"{verdict.decision.value}"
            )

        reports = build_reports(
            job_id="demo",
            state=job.state,
            ledger=__import__(
                "oatutor_council.persistence", fromlist=["load_ledger"]
            ).load_ledger(db, "demo"),
            changes=changes,
            verdicts=list_verdicts(db, "demo"),
            attempts=__import__(
                "oatutor_council.persistence", fromlist=["list_attempts"]
            ).list_attempts(db, "demo"),
            # Read back from the rows the job wrote, exactly as `GET /report` does.
            findings=latest_findings(db, "demo", kind="content"),
            integrity_findings=latest_findings(db, "demo", kind="integrity"),
            usage=token_usage(db, "demo"),
            artifacts=describe_artifacts(db, "demo"),
            rediscoveries=rediscovery_counts(db, "demo"),
        )

        print("\n--- final report --------------------------------------------------")
        print(render_markdown(reports))

        print("--- guarantees ----------------------------------------------------")
        print(f"  source workbook unchanged: {sha256_of(source) == source_hash}")
        corrected = read_workbook(copy.path)
        print(f"  corrected workbook reopens: {len(corrected.blocks)} block(s)")
        print(f"  final state: {job.state.value}")
        print(f"  model calls: {job.llm_calls_used}, steps: {job.steps_used}")

        # Read from `llm_calls`, not from the mock's memory. The isolation claim is
        # about what was transmitted, and the rows are the record of that.
        reviewer_payloads = [
            call["payload"]["user_payload"]
            for call in list_llm_calls(db, "demo")
            if call["role"]
            in (
                AgentRole.KNOWN_ISSUE_REVIEWER.value,
                AgentRole.INDEPENDENT_REVIEWER.value,
            )
        ]
        leaked = [p for p in reviewer_payloads if "Excel turned the fraction" in p]
        print(f"  reviewer payloads inspected: {len(reviewer_payloads)}, leaks: {len(leaked)}")

        if args.keep:
            print(f"\nworkspace kept at {workspace}")
        return 0
    finally:
        if not args.keep:
            shutil.rmtree(workspace, ignore_errors=True)


if __name__ == "__main__":
    raise SystemExit(main())
